"""Build the editable Word rendition from the single Markdown source.
Usage: bundled python build_resume.py. Then run export_pdf.ps1.
"""
from pathlib import Path
from datetime import datetime
import hashlib
import json
import re
import shutil
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path(__file__).resolve().parent.parent
ROOT = BASE.parent
SOURCE = BASE / 'source/이력서.md'
OUT = BASE / 'output'
STEM = '박종혁_기술영업_이력서'
OUT.mkdir(exist_ok=True)
profile = json.loads((ROOT / '공통자료/profile.json').read_text(encoding='utf-8-sig'))
source = SOURCE.read_text(encoding='utf-8-sig')
text = source.replace('{{연락처}}', profile['phone']).replace('{{이메일}}', profile['email'])
if '{{' in text:
    raise ValueError('Unresolved placeholder in resume')

previous = [p for p in [OUT / (STEM+'.docx'), OUT / (STEM+'.pdf')] if p.exists()]
if previous:
    history = BASE / 'versions' / datetime.now().strftime('%Y-%m-%d_%H%M%S')
    history.mkdir(parents=True, exist_ok=False)
    for path in previous:
        shutil.copy2(path, history / path.name)
    for path, name in [
        (BASE/'.qa/latest_source_snapshot.md','이력서.md'),
        (BASE/'source/build_manifest.json','build_manifest.json')
    ]:
        if path.exists():shutil.copy2(path,history/name)
    (history / '출처.json').write_text(json.dumps({
        'role': '교체 전 출력 보존본. 현행본 아님',
        'files': {p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in previous}
    }, ensure_ascii=False, indent=2), encoding='utf-8')

doc = Document()
sec = doc.sections[0]
sec.page_width = Mm(210)
sec.page_height = Mm(297)
sec.top_margin = Mm(16)
sec.bottom_margin = Mm(16)
sec.left_margin = Mm(18)
sec.right_margin = Mm(18)
sec.header_distance = Mm(7)
sec.footer_distance = Mm(8)

FONT = '맑은 고딕'
def font_style(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0,0,0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    for key in ('ascii','hAnsi','eastAsia','cs'):
        fonts.set(qn('w:'+key), FONT)
    for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        fonts.attrib.pop(qn('w:'+key),None)

normal = doc.styles['Normal']
font_style(normal, 10.5)
normal.paragraph_format.line_spacing = Pt(15.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
for name, size, before, after in [
    ('Title',25,0,4), ('Subtitle',11,0,4),
    ('Heading 1',12.5,11,5), ('Heading 2',11,7,3),
]:
    st = doc.styles[name]
    font_style(st,size,name!='Subtitle')
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.line_spacing = 1.15
    st.paragraph_format.keep_with_next = True
    ppr=st.element.find(qn('w:pPr'))
    if ppr is not None:
        for tag in ('pBdr','shd'):
            found=ppr.find(qn('w:'+tag))
            if found is not None:ppr.remove(found)

font_style(doc.styles['Header'],8)
font_style(doc.styles['Footer'],8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('박종혁  |  ')
for field in ['PAGE','NUMPAGES']:
    if field=='NUMPAGES':footer.add_run(' / ')
    el=OxmlElement('w:fldSimple'); el.set(qn('w:instr'),field)
    footer._p.append(el)

def rich(p, value):
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',value)):
        if part:p.add_run(part).bold=(i%2==1)

blocks = [b.strip() for b in re.split(r'\n\s*\n',text) if b.strip()]
after_title=False
for block in blocks:
    if block=='<!-- pagebreak -->':
        doc.add_page_break()
        continue
    if block.startswith('# '):
        doc.add_paragraph(block[2:],'Title'); after_title=True
    elif block.startswith('### '):
        title=block[4:]
        title=re.sub(r'^(\d+)\.\s*',r'\1 ',title).replace(' | ',' ')
        doc.add_paragraph(title,'Heading 2')
    elif block.startswith('## '):
        doc.add_paragraph(block[3:].replace('·',' '),'Heading 1')
    elif block.startswith('- '):
        for line in block.splitlines():
            if not line.startswith('- '):raise ValueError('Unsupported bullet block')
            p=doc.add_paragraph()
            p.paragraph_format.left_indent=Mm(3)
            p.paragraph_format.first_line_indent=Mm(-3)
            p.paragraph_format.keep_together=True
            p.add_run('• ')
            rich(p,line[2:])
    elif after_title:
        doc.add_paragraph(block,'Subtitle'); after_title=False
    else:
        p=doc.add_paragraph()
        rich(p,block.replace('\n',' '))
        if re.match(r'^202[0-9]\.',block):
            p.paragraph_format.space_after=Pt(4)
            p.paragraph_format.keep_with_next=True
            for run in p.runs:run.font.size=Pt(9)
        if block.startswith(profile['phone']):
            p.paragraph_format.space_after=Pt(5)
            for run in p.runs:run.font.size=Pt(9)

doc.core_properties.title='박종혁 기술영업 이력서'
doc.core_properties.subject='업무용 솔루션 기술영업 지원'
doc.core_properties.author='박종혁'
doc.core_properties.keywords=''
doc.save(OUT/(STEM+'.docx'))
meta={
    'built_at':datetime.now().isoformat(timespec='seconds'),
    'source':'source/이력서.md',
    'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
    'facts':'../공통자료/career_facts.md',
    'facts_sha256':hashlib.sha256((ROOT/'공통자료/career_facts.md').read_bytes()).hexdigest(),
    'contact_values_logged':False,
    'outputs':[STEM+'.docx',STEM+'.pdf'],
    'pdf_status':'pending_export',
}
(BASE/'source/build_manifest.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2),encoding='utf-8')
(BASE/'.qa').mkdir(exist_ok=True)
(BASE/'.qa/latest_source_snapshot.md').write_text(source,encoding='utf-8')
print('Word generated from Markdown; contact placeholders resolved')
