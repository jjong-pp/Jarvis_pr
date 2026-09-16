# -*- coding: utf-8 -*-
"""박종혁 PM 포트폴리오 v3 — 22장 벡터

v1.2 대비 변경
1. 헤더 부제(소제목 옆 한 줄 설명) 전면 삭제
2. 하단 풀폭 알약 바(마무리 한 줄) 전면 삭제 — 22장 중 12장이 동일 수사 골격이었다
3. 비워진 자리를 내용으로 채우고, 결론은 문장이 아니라 표·수치로 남긴다
4. 통이미지 → 벡터. 텍스트 추출이 되므로 ATS 파싱 문제도 함께 해소된다

잔여 수정 반영 (career_facts 10절 · 2차 메타 감사)
- 1장  손익분기까지 계산 → 대시보드 근거로 교체 / 현업 인터뷰·단위경제 계산 삭제
- 7장  10:00 운송장 배정·제출 → 11:10 처리창 시작
- 8장  익일 재처리 행 삭제 / 자동 루프 → 처리창 내 + 미해결 시 수동 통관 격리
- 11장 네 기준 모두 충족 → 장점란 근거에 맞게
- 14장 정산 TO-BE에 1명 복원 / 420분 기준 명시 / 오출고 각주
- 15장 사내 크롤링 로직·이미 있던 자산 → 표현 조정

2026-09-06 원장 갱신 반영
- C-081 `2개월 미사용`은 폐기. `2026.04 시범 배포 → 약 2개월 병행 검증 → 2026.06 완전 이관`
"""
import os
from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "박종혁_PM_포트폴리오_v3.pptx")
PHOTO = os.path.abspath(os.path.join(HERE, "..", "..", "assets", "portrait", "_photo.png"))

F = "Malgun Gothic"
INK, SUB, MUTE = "1A1A1A", "6E6E73", "9A9AA0"
LINE, SOFT, WHITE = "D6D6DB", "F4F4F6", "FFFFFF"
RED, GOLD, GOLD_BG = "B42318", "9A6400", "F3E5CC"
GRN, GRN_BG = "1E7A46", "E7F5EC"

PAL = {
    "kcs": dict(pri="1B4F9C", dk="16437F", br="3B72C4", bg="E9F1FC", bg2="F1F6FC", mid="8FB2E0", lt="D2DFF1"),
    "b2b": dict(pri="1E6B45", dk="17563A", br="3E9268", bg="E8F5EE", bg2="F0F9F4", mid="9ED4B6", lt="CDEAD9"),
    "mb":  dict(pri="B34700", dk="8F3900", br="E06C1F", bg="FFF6E6", bg2="FDF7EF", mid="E0A878", lt="F7D9BE"),
    "nu":  dict(pri="2A3F6B", dk="1E2E4F", br="4A648F", bg="EFF2F7", bg2="F5F7FA", mid="B9C4D8", lt="D9E0EC"),
}

ML, CW, RGT = 0.667, 11.995, 12.662
COL_W, COL_L, COL_R = 5.830, 0.667, 6.832
BODY = 1.560          # 부제를 없앤 만큼 본문 시작선이 올라온다


def C(h): return RGBColor.from_string(h)


def blank(pr): return pr.slides.add_slide(pr.slide_layouts[6])


def tb(sl, x, y, w, h, size=11, color=INK, bold=False,
       align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    s = sl.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    pr = tf.paragraphs[0]
    pr.alignment = align
    pr.line_spacing = spacing
    s._d = (size, color, bold)
    return s


def put(sh, t, size=None, color=None, bold=None, spacing=None, after=0):
    tf = sh.text_frame
    first = (len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs)
    pg = tf.paragraphs[0] if first else tf.add_paragraph()
    pg.alignment = tf.paragraphs[0].alignment
    if spacing is not None: pg.line_spacing = spacing
    pg.space_after = Pt(after)
    r = pg.add_run(); r.text = t
    ds, dc, db = sh._d
    r.font.name = F
    r.font.size = Pt(size if size is not None else ds)
    r.font.bold = db if bold is None else bold
    r.font.color.rgb = C(color if color is not None else dc)
    return pg


def text(sl, x, y, w, h, t, size=11, color=INK, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    sh = tb(sl, x, y, w, h, size, color, bold, align, anchor, spacing)
    put(sh, t)
    return sh


def rect(sl, x, y, w, h, fill=None, line=None, lw=0.75,
         shape=MSO_SHAPE.RECTANGLE, radius=None):
    s = sl.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: s.adjustments[0] = min(0.5, radius / min(w, h))
        except Exception: pass
    if fill: s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    else: s.fill.background()
    if line: s.line.color.rgb = C(line); s.line.width = Pt(lw)
    else: s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(0.10)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return s


def rr(sl, x, y, w, h, fill=None, line=None, radius=0.07):
    return rect(sl, x, y, w, h, fill, line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)


def label(sl, x, y, w, h, t, size=10, color=WHITE, fill=None, line=None,
          bold=True, align=PP_ALIGN.CENTER, radius=None):
    sh = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = rect(sl, x, y, w, h, fill, line, shape=sh, radius=radius)
    pg = s.text_frame.paragraphs[0]; pg.alignment = align
    r = pg.add_run(); r.text = t
    r.font.name = F; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = C(color)
    return s


def hline(sl, x, y, w, color=LINE, lw=1.0):
    s = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, I(x), I(y), I(w), Pt(lw))
    s.fill.solid(); s.fill.fore_color.rgb = C(color)
    s.line.fill.background(); s.shadow.inherit = False
    return s


def head(sl, num, title, page, pal):
    """부제 없음. 번호 + 제목만."""
    text(sl, 0.682, 0.855, 0.480, 0.300, num, 21, pal["pri"], True, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, 1.230, 0.855, 8.500, 0.300, title, 21, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(sl, 10.462, 0.345, 2.200, 0.260, page, 10.5, MUTE, False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    hline(sl, ML, 1.290, CW)


def cap(sl, x, y, w, t, color=None):
    return text(sl, x, y, w, 0.240, t, 10, color or SUB, True)


def table(sl, x, y, cols, rows, rowh=0.520, zebra=None, hi=None, hifill=None,
          pal=None, sz=11.2, hsz=9.8):
    """cols=[(라벨,폭)...] rows=[[셀...]...] hi=강조할 행 인덱스"""
    cx = x
    for lb, w in cols:
        text(sl, cx, y, w, 0.240, lb, hsz, SUB, True); cx += w
    tw = sum(w for _, w in cols)
    hline(sl, x, y + 0.300, tw, pal["mid"] if pal else LINE)
    ry = y + 0.400
    for i, r in enumerate(rows):
        if hi is not None and i == hi:
            rect(sl, x, ry - 0.045, tw, rowh, hifill or (pal["bg"] if pal else SOFT))
        elif zebra and i % 2 == 0:
            rect(sl, x, ry - 0.045, tw, rowh, zebra)
        cx = x
        for j, (lb, w) in enumerate(cols):
            text(sl, cx, ry, w - 0.150, rowh - 0.090, r[j], sz,
                 INK if j == 0 else SUB, j == 0,
                 anchor=MSO_ANCHOR.MIDDLE, spacing=1.22)
            cx += w
        ry += rowh
    hline(sl, x, ry - 0.045, tw, pal["mid"] if pal else LINE)
    return ry


def cover(sl, pal, page, ptitle, badge, badge_sub, headline, lede, fields,
          panel_title, panel_rows, meta):
    rect(sl, -0.2, -0.2, 13.8, 8.0, pal["pri"])
    text(sl, 10.462, 0.345, 2.200, 0.260, page, 10.5, pal["lt"], False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
    text(sl, ML, 0.840, 11.500, 0.320, ptitle, 15, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    rr(sl, ML, 1.750, 0.950, 0.953, WHITE, radius=0.10)
    text(sl, ML, 1.750, 0.950, 0.620, badge, 20, pal["pri"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(sl, ML, 2.360, 0.950, 0.300, badge_sub, 8, pal["pri"], False, PP_ALIGN.CENTER)
    text(sl, 1.900, 1.780, 6.450, 0.380, headline, 21.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
    ld = tb(sl, 1.900, 2.230, 6.450, 0.620, 13, pal["lt"], spacing=1.48)
    for l in lede: put(ld, l)
    for ly, k, vy, lines in fields:
        text(sl, 1.900, ly, 6.450, 0.250, k, 12, pal["mid"], True, anchor=MSO_ANCHOR.MIDDLE)
        b = tb(sl, 1.900, vy, 6.550, 0.520, 12.8, WHITE, spacing=1.45)
        for l in lines: put(b, l)
    rr(sl, 8.458, 1.750, 4.203, 3.391, pal["dk"], radius=0.11)
    text(sl, 8.758, 1.980, 3.600, 0.260, panel_title, 12, pal["lt"], True, anchor=MSO_ANCHOR.MIDDLE)
    ry = 2.430
    for a, b in panel_rows:
        rr(sl, 8.758, ry, 3.603, 0.470, pal["pri"])
        text(sl, 8.958, ry, 2.150, 0.470, a, 12.5, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 11.111, ry, 1.050, 0.470, b, 11, pal["lt"], False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
        ry += 0.640
    hline(sl, ML, 5.470, CW, pal["mid"], 0.9)
    my = 5.840
    for k, v in meta:
        text(sl, ML, my, 0.900, 0.260, k, 10.5, pal["mid"], True, anchor=MSO_ANCHOR.MIDDLE)
        text(sl, 1.600, my, RGT - 1.600, 0.260, v, 10.2, pal["bg2"], anchor=MSO_ANCHOR.MIDDLE)
        my += 0.280


p = Presentation()
p.slide_width, p.slide_height = I(13.3333), I(7.5)
K, B, M, N = PAL["kcs"], PAL["b2b"], PAL["mb"], PAL["nu"]

# ═════════════════════════════════════════ 01 표지·자기소개
s = blank(p)
rect(s, 0, 0, 13.3333, 7.5, WHITE)
rect(s, 8.900, 0, 4.4333, 7.5, SOFT)
text(s, ML, 1.020, 7.600, 1.480,
     "현장에서 겪은 운영 문제를\n정책과 데이터 구조로 바꿔\n실제로 돌아가게 만듭니다.",
     26, INK, True, spacing=1.32)
bd = tb(s, ML, 2.780, 7.600, 1.600, 12.8, SUB, spacing=1.58)
put(bd, "아이베(EIBE) SCM팀에서 수입·물류 실무를 하며 반복되는 병목을 직접 겪었고,")
put(bd, "그 문제를 요구사항과 데이터 구조로 바꿔 외주 개발사·현업과 함께 시스템으로 만들었습니다.")
put(bd, "규제 대응, 사업 확장에 따른 운영 디지털화, 수요예측 도구, SaaS 제품 기획까지 담당했습니다.")

cap(s, ML, 4.330, 3.000, "강점")
for i, (t1, t2, ev, key) in enumerate([
        ("현장 구조화", "수기 업무와 계약 예외를\n요구사항으로 바꿉니다", "B2B · 계약 예외를 설정값으로", "b2b"),
        ("이해관계 조율", "현업·개발사·공급사의\n서로 다른 기준을 맞춥니다", "관세청 · 4개 주체의 책임 분리", "kcs"),
        ("운영 안착", "배포가 아니라 현업이 쓰고\n숫자가 바뀌는 것까지 봅니다", "대시보드 · 2개월 병행 검증 후 이관", "mb")]):
    x = ML + i * 2.720
    text(s, x, 4.640, 0.360, 0.280, "0%d" % (i + 1), 12, PAL[key]["pri"], True)
    text(s, x, 4.990, 2.520, 0.300, t1, 15, INK, True)
    text(s, x, 5.360, 2.560, 0.620, t2, 11.3, SUB, spacing=1.32)
    text(s, x, 6.060, 2.560, 0.280, ev, 10.2, PAL[key]["pri"], True)

hline(s, ML, 6.560, 7.600)
cap(s, ML, 6.690, 3.000, "다루는 프로젝트")
for i, (nm, ds, key) in enumerate([("규제 대응", "관세청 통관 API 연동", "kcs"),
                                   ("운영 디지털화", "B2B 산후조리원 전용몰", "b2b"),
                                   ("제품 기획", "마켓빅시 BIGSEE 리서치 SaaS", "mb")]):
    x = ML + i * 2.720
    rect(s, x, 6.980, 0.070, 0.340, PAL[key]["pri"])
    text(s, x + 0.190, 6.960, 2.480, 0.260, nm, 10.8, PAL[key]["pri"], True)
    text(s, x + 0.190, 7.190, 2.480, 0.240, ds, 10, SUB)

if os.path.exists(PHOTO):
    s.shapes.add_picture(PHOTO, I(9.560), I(0.700), height=I(1.900))
PX = 9.320
cap(s, PX, 2.870, 3.400, "PROFILE")
text(s, PX, 3.130, 3.600, 0.290, "박종혁", 17, INK, True)
text(s, PX, 3.470, 3.600, 0.240, "010-7449-6865", 10.3, SUB)
text(s, PX, 3.700, 3.600, 0.240, "parkjonghyeok2000@gmail.com", 10.3, SUB)
hline(s, PX, 4.010, 3.500)
cap(s, PX, 4.120, 3.400, "CAREER")
text(s, PX, 4.380, 3.600, 0.250, "아이베(EIBE) · SCM팀", 12, INK, True)
text(s, PX, 4.640, 3.600, 0.240, "2025.09 ~ 재직 중", 10.3, SUB)
text(s, PX, 4.870, 3.600, 0.240, "SCM 실무 · 사내 IT 기획 병행", 10.3, SUB)
hline(s, PX, 5.170, 3.500)
cap(s, PX, 5.280, 3.400, "EDUCATION")
text(s, PX, 5.540, 3.600, 0.250, "고려사이버대 AI·데이터과학부", 11.3, INK, True)
text(s, PX, 5.790, 3.600, 0.240, "2025.02 편입 · 4.1 / 4.5", 10.3, SUB)
text(s, PX, 6.030, 3.600, 0.250, "부천대 컴퓨터소프트웨어과", 11.3, INK, True)
text(s, PX, 6.280, 3.600, 0.240, "2019.03 ~ 2024.02 · 3.8 / 4.5", 10.3, SUB)
hline(s, PX, 6.580, 3.500)
cap(s, PX, 6.690, 3.400, "SKILL")
text(s, PX, 6.950, 3.700, 0.240, "요구사항·정책 설계 · 데이터 구조 · QA/UAT", 9.8, SUB)
text(s, PX, 7.170, 3.700, 0.240, "Python · FastAPI · SQLAlchemy · SQL · pandas · Git", 9.8, SUB)

# ═════════════════════════════════════════ 02 목차
s = blank(p)
rect(s, 0, 0, 13.3333, 7.5, WHITE)
text(s, 10.462, 0.345, 2.200, 0.260, "02 / 22", 10.5, MUTE, False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
text(s, ML, 3.150, 3.000, 0.280, "Portfolio", 13, SUB)
text(s, ML, 3.480, 4.000, 0.800, "Project", 46, INK, True)
LX, CXN, CXB = 5.850, 4.930, 6.270
rect(s, LX, 1.760, 0.022, 4.760, LINE)
for i, (num, key, cat, tag, title, m1, m2) in enumerate([
    ("01", "kcs", "규제 대응", "강점 · 이해관계 조율", "관세청 사전 거래정보 신고(TRA) 연동",
     "p 03 ~ 08  |  전체 구매 여정 및 로직 기획",
     "2026.02 ~ 08 · 개발비 약 4,000만원 · 데이터 4개 주체 · 조율 5개사"),
    ("02", "b2b", "운영 디지털화", "강점 · 현장 구조화", "B2B 산후조리원 전용몰 구축",
     "p 09 ~ 14  |  기획·QA 총괄",
     "2026.06 기획 ~ 08.31 최종 검수 · 계약처 약 140개소 · 공급사 5곳"),
    ("03", "mb", "제품 기획", "강점 · 운영 안착", "마켓빅시 BIGSEE 리서치 SaaS",
     "p 15 ~ 21  |  제품·사업 기획 (개발 구현은 동료 전담)",
     "진행 중 · 등록 사용자 약 20명 (활성·유료 미확인)")]):
    col = PAL[key]["pri"]; y = 1.760 + i * 1.640
    text(s, CXN, y - 0.180, 1.000, 0.700, num, 40, "E7E7EA", True, PP_ALIGN.RIGHT)
    rect(s, LX - 0.058, y + 0.055, 0.138, 0.138, col, shape=MSO_SHAPE.OVAL)
    text(s, CXB, y - 0.020, 2.400, 0.260, cat, 11.5, col, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, CXB + 2.500, y - 0.020, 3.000, 0.260, tag, 10.5, SUB, anchor=MSO_ANCHOR.MIDDLE)
    text(s, CXB, y + 0.290, 6.400, 0.380, title, 21, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, CXB, y + 0.740, 6.400, 0.260, m1, 11.5, SUB, anchor=MSO_ANCHOR.MIDDLE)
    text(s, CXB, y + 1.030, 6.400, 0.260, m2, 10.5, MUTE, anchor=MSO_ANCHOR.MIDDLE)
text(s, ML, 6.780, CW, 0.280,
     "성과 수치는 각 프로젝트 장에서 측정 기간·모수·원천과 함께 제시합니다.", 10, SUB)

# ═════════════════════════════════════════ 03 관세청 표지
s = blank(p)
cover(s, K, "01 / 06", "Project 01.   규제 대응 — 주문부터 출고까지 전사 흐름 재설계",
      "KCS", "TRA · PER API",
      "네 주체가 나눠 갖던 데이터를 하나의 신고로",
      ["주문 한 건의 데이터가 자사몰·본인확인기관·해외 ERP·통관대행사에 흩어져",
       "어디도 신고 한 건을 혼자 완성할 수 없었습니다."],
      [(2.990, "Concept", 3.290, ["신고 판단과 상태 통제의 중심을 해외 ERP가 아닌 국내 DB로 이전"]),
       (3.650, "Main Target", 3.900, ["자사몰·외부 판매채널에서 해외직구를 구매하는 고객과 운영팀"]),
       (4.310, "Project Goal", 4.560, ["미신고 출고 0건 · 채널 확장에 견디는 제출 구조 · 오류 복구 가능한 운영"])],
      "데이터를 나눠 가진 네 주체",
      [("자사몰", "주문·결제·고객"), ("본인확인기관", "통관부호 검증"),
       ("해외 ERP·WMS", "상품·운송장"), ("통관대행사", "수입신고·병합")],
      [("기간", "2026.02 ~ 2026.08"),
       ("Role", "정책·API 해석 / 대안 비교 / 요구사항·데이터 계약 정의 / 관계사 조율 / QA / 운영정책 수립"),
       ("협업", "자사몰 개발사 · ERP 개발사 · 물류 개발사 · 통관대행사 · 본인확인기관 (역할명 표기)"),
       ("산출물", "통합 기획서 · API 필드 매핑표 · 리스크/질의 대장 · 정책 변경 이력 · E2E 시퀀스 · QA 시나리오"),
       ("예산", "개발비 약 4,000만원")])

# ═════════════════════════════════════════ 04 배경과 문제 정의
s = blank(p); head(s, "01", "배경과 문제 정의", "02 / 06", K)
rect(s, ML, BODY, CW, 1.180, SOFT)
bg = tb(s, ML + 0.280, BODY + 0.150, CW - 0.560, 0.900, 11.2, INK, spacing=1.42)
put(bg, "1.   관세법 제254조가 2024.12.31 개정되어 2026.01.01 시행됩니다. 간이통관을 이용하려면 관세청에 등록해야 합니다.")
put(bg, "2.   시행령 제258조는 주문번호·구매일자·수하인 성명·PCCC·품명·수량·결제금액을 결제 완료 후 수입 전까지 전자적으로 제공하도록 구체화했습니다.")
put(bg, "3.   전자상거래 통관플랫폼에서는 전자상거래업자부호가 신고서·통관목록의 필수값이 되고, PCCC 도용방지와 공급망 단위 위험관리가 적용됩니다.")

cap(s, COL_L, 3.030, COL_W, "AS-IS 1 — 데이터와 책임이 회사별로 나뉘어 있었다", K["pri"])
for i, (a, b, loc) in enumerate([("자사몰", "주문 · 결제 · 고객 정보", "국내"),
                                 ("본인확인기관", "개인통관고유부호 유효성 검증", "국내"),
                                 ("ERP · WMS", "상품 · 포장 · 운송장 · 실중량", "해외"),
                                 ("통관대행사", "수입신고 · 후속 신고", "국내")]):
    y = 3.380 + i * 0.560
    rr(s, COL_L, y, COL_W, 0.470, WHITE, LINE)
    text(s, COL_L + 0.220, y, 1.700, 0.470, a, 11.5, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 1.960, y, 2.900, 0.470, b, 10.5, SUB, anchor=MSO_ANCHOR.MIDDLE)
    label(s, COL_L + 4.950, y + 0.095, 0.680, 0.280, loc, 9.5,
          RED if loc == "해외" else K["pri"], "FCEAEA" if loc == "해외" else K["bg"], radius=0.14)

cap(s, COL_R, 3.030, COL_W, "AS-IS 2 — 기존 Excel 수기 신고의 한계", K["pri"])
for i, t in enumerate(["제출 상태 · 오류 · 재처리 이력을 시스템으로 추적할 수 없습니다.",
                       "판매 채널이 늘어날수록 수작업과 중복 제출 위험이 함께 커집니다.",
                       "누가 언제 무엇을 제출했는지 남지 않아 사후 감사에 대응하기 어렵습니다."]):
    y = 3.400 + i * 0.480
    rect(s, COL_R, y + 0.140, 0.100, 0.100, K["mid"], shape=MSO_SHAPE.OVAL)
    text(s, COL_R + 0.240, y, COL_W - 0.240, 0.400, t, 11.2, INK, spacing=1.25)
rr(s, COL_R, 4.900, COL_W, 0.540, GOLD_BG)
text(s, COL_R + 0.240, 4.900, COL_W - 0.480, 0.540,
     "비상 대응 수단으로는 유지하되 상시 운영 방식에서는 제외", 11, GOLD, True, anchor=MSO_ANCHOR.MIDDLE)

rect(s, ML, 5.760, CW, 1.140, K["bg"])
text(s, ML + 0.320, 5.900, 6.400, 0.860,
     "네 곳 중 어디도 신고 한 건을 혼자 완성할 수 없었습니다.\n누가 어떤 데이터를 책임지고, 어디서 하나의 신고로 완성할 것인가.",
     13.5, K["dk"], True, spacing=1.42)
text(s, 7.500, 5.940, 5.160, 0.780,
     "기능 요구가 아니라 데이터 소유권과 책임의 문제로 정의한 것이\n이 프로젝트의 출발점이었고, 이후 모든 대안 비교의 판단 기준이 됐습니다.",
     10.8, SUB, spacing=1.35)

# ═════════════════════════════════════════ 05 대안 비교와 책임 분담
s = blank(p); head(s, "02", "대안 비교와 책임 분담", "03 / 06", K)
cap(s, ML, BODY, 5.000, "검토한 세 가지 구조", K["pri"])
ry = table(s, ML, BODY + 0.310,
           [("대안", 2.800), ("장점", 2.500), ("한계", 3.600), ("판단 근거", 1.700), ("결정", 1.395)],
           [["해외 ERP 개발사 중심 통제", "상품·출고 원천 데이터 활용이 쉬움",
             "개인정보 국외 이전 우려 · 해외 개발조직의 API 개방 승인·권한 확보 난이도", "개인정보 보호", "제외"],
            ["Excel 수기 제출 유지", "즉시 대응 가능 · 추가 개발 없음",
             "제출 상태·재처리·감사 추적 불가 · 채널 확장 시 중복 제출 위험", "운영 추적성", "장애 시 폴백"],
            ["국내 자사몰 DB 중심 통제", "개인정보 최소화 · 제출 상태 통제 · 원천별 책임 분리",
             "신규 인터페이스와 회사 간 데이터 계약이 필요", "세 기준 모두 충족", "채택"]],
           rowh=0.620, hi=2, pal=K, sz=10.8)
rect(s, ML, ry + 0.020, CW, 0.400, SOFT)
text(s, ML + 0.240, ry + 0.020, CW - 0.480, 0.400,
     "판단 기준 순서   ①  개인정보 최소화      ②  제출 상태 통제 가능성      ③  원천별 책임 분리      ④  개발 난이도",
     11, INK, anchor=MSO_ANCHOR.MIDDLE)

cap(s, ML, 4.680, 5.000, "결정 후 각 주체가 책임지는 범위", K["pri"])
text(s, 7.400, 4.680, 5.162, 0.240,
     "R 실행 책임  ·  A 최종 승인  ·  C 협의  ·  I 통보", 10, MUTE, False, PP_ALIGN.RIGHT)
rcols = [("역할 / 데이터", 2.800), ("자사몰", 1.150), ("본인확인기관", 1.500),
         ("ERP · WMS", 1.300), ("통관대행사", 1.400), ("비고", 3.845)]
cx = ML
for j, (lb, w) in enumerate(rcols):
    text(s, cx, 5.000, w, 0.240, lb, 9.8, SUB, True,
         PP_ALIGN.CENTER if 1 <= j <= 4 else PP_ALIGN.LEFT)
    cx += w
hline(s, ML, 5.300, CW, K["mid"])
ry = 5.400
for r in [["주문 · 결제 · 고객 정보", "R", "·", "·", "I", "원천 보유 · 제출 주체"],
          ["개인통관고유부호 검증", "A", "R", "·", "·", "국내에서만 처리"],
          ["상품 · 포장 · 운송장 · 실중량", "I", "·", "R", "C", "해외 시스템이 원천 책임"],
          ["관세청 제출과 상태 관리", "R", "·", "C", "I", "통제점을 국내로 이전"],
          ["운송장 기준 병합 · 후속 신고", "C", "·", "C", "R", "주문번호 + 운송장으로 병합"]]:
    cx = ML
    for j, (lb, w) in enumerate(rcols):
        v = r[j]
        if 1 <= j <= 4:
            bx = cx + (w - 0.400) / 2          # 헤더 라벨과 같은 축에 배지를 놓는다
            if v == "R":
                label(s, bx, ry + 0.025, 0.400, 0.320, "R", 10, WHITE, K["pri"], radius=0.05)
            elif v in ("A", "C", "I"):
                label(s, bx, ry + 0.025, 0.400, 0.320, v, 10, K["dk"], K["lt"], radius=0.05)
            else:
                text(s, cx, ry, w, 0.370, "·", 11, MUTE, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        else:
            text(s, cx, ry, w - 0.150, 0.370, v, 10.8, INK if j == 0 else SUB, j == 0, anchor=MSO_ANCHOR.MIDDLE)
        cx += w
    ry += 0.380
hline(s, ML, ry - 0.020, CW, K["mid"])

# ═════════════════════════════════════════ 06 E2E 흐름
s = blank(p); head(s, "03", "E2E 흐름과 사용자 시나리오", "04 / 06", K)
LW, LX0 = 1.720, 2.180
for i, ln in enumerate(["고객", "자사몰", "본인확인기관", "ERP · WMS", "관세청", "통관대행사"]):
    rect(s, LX0 + i * (LW + 0.060), BODY, LW, 0.380, INK)
    text(s, LX0 + i * (LW + 0.060), BODY, LW, 0.380, ln, 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
y = BODY + 0.470
for st, nm, cells, note in [
    ("Stage 0", "최초 1회 본인확인", [(0, "본인확인 요청"), (1, "인증키 발급 요청"), (2, "인증키 반환")],
     "인증키는 국내 회원 DB에 암호화 저장 · 원문 미저장"),
    ("Stage 1", "결제 시점 검증", [(0, "통관부호 입력·결제"), (1, "간소화 검증 요청"), (2, "검증 결과 회신")],
     "불일치·무효면 결제 차단 후 수정 또는 일반 인증으로 전환"),
    ("Stage 2", "주문 수집", [(1, "결제완료 주문 확정"), (3, "주문 수집 · 상태 변경")],
     "이 시점부터 고객 취소·환불 버튼 제거"),
    ("Stage 3", "운송장 배정·제출", [(1, "전건 배정 확인 후 제출"), (3, "운송장·포장 확정"), (4, "신고 접수")],
     "일부만 배정되면 대기 상태 유지 · 부분 제출 차단"),
    ("Stage 4", "오류 확인·복구", [(1, "원천 구분 후 수정·재제출"), (4, "오류 코드 회신")],
     "약 1분 내 회신 · 관리자 화면에서 원인 확인"),
    ("Stage 5", "병합·후속 신고", [(3, "실물 물류 확정 전송"), (4, "수입신고"), (5, "운송장+주문번호로 병합")],
     "미신고 상태에서는 출고되지 않도록 차단")]:
    rect(s, ML, y, 1.420, 0.560, K["bg"])
    text(s, ML, y + 0.040, 1.420, 0.250, st, 10.5, K["dk"], True, PP_ALIGN.CENTER)
    text(s, ML, y + 0.280, 1.420, 0.250, nm, 9.2, SUB, False, PP_ALIGN.CENTER)
    for ci, ct in cells:
        cx = LX0 + ci * (LW + 0.060)
        rr(s, cx, y, LW, 0.400, WHITE, LINE)
        text(s, cx + 0.070, y, LW - 0.140, 0.400, ct, 9.6, INK, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.12)
    text(s, LX0, y + 0.410, 8.000, 0.220, note, 9.4, MUTE)
    y += 0.720
rect(s, ML, y + 0.030, 5.400, 0.560, INK)
text(s, ML + 0.240, y + 0.030, 5.160, 0.560,
     "운송장 1건  =  해외 포장 1건  =  제출 1건  =  통관 신고 1건", 12, WHITE, True, anchor=MSO_ANCHOR.MIDDLE)
text(s, 6.400, y + 0.070, 6.262, 0.480,
     "회사마다 다르게 부르던 단계를 하나의 그림으로 맞춰,\n개발 범위 협의의 기준 문서가 됐습니다.", 10.8, SUB, spacing=1.35)

# ═════════════════════════════════════════ 07 요구사항과 산출물
s = blank(p); head(s, "04", "요구사항과 실제 산출물", "05 / 06", K)
cap(s, ML, BODY, 7.000, "운영 정책 변경 — 당일 출고 주문 마감 11:00 → 09:00", K["pri"])
TW = (CW - 0.300) / 6
for i, (t, d) in enumerate([("09:00", "주문 수집 마감"), ("11:10", "처리창 시작"), ("11:40", "오류 조회"),
                            ("11:50", "처리창 마감 · 미해결 격리"), ("12:00", "출고 취소 마감"), ("13:00", "출고")]):
    x = ML + i * (TW + 0.060)
    rr(s, x, BODY + 0.330, TW, 0.640, K["pri"] if i == 0 else WHITE, None if i == 0 else LINE)
    text(s, x, BODY + 0.400, TW, 0.280, t, 13, WHITE if i == 0 else K["dk"], True, PP_ALIGN.CENTER)
    text(s, x + 0.060, BODY + 0.690, TW - 0.120, 0.260, d, 9.2,
         K["lt"] if i == 0 else SUB, False, PP_ALIGN.CENTER)
text(s, ML, BODY + 1.070, CW, 0.260,
     "마감을 2시간 앞당기는 대신 주문 안내 문구·배너·CS 응대 기준을 함께 고쳤습니다.", 10.8, SUB)

cap(s, COL_L, 3.320, COL_W, "데이터 계약 — 필드가 아니라 값의 의미를 고정했다", K["pri"])
for i, (a, b) in enumerate([("고객 최종 결제액", "운송장별 신고가액과 분리해 관리"),
                            ("원주문번호", "ERP 분할번호 · 운송장을 별도 식별자로 유지"),
                            ("인증 정보", "해외에 전달하지 않고 국내에서 통관대행사로 직접 전달"),
                            ("상품 · 실중량", "해외 ERP·WMS가 책임지고 자사몰은 재계산하지 않음")]):
    y = 3.680 + i * 0.660
    rect(s, COL_L, y, 1.700, 0.580, K["bg2"])
    text(s, COL_L + 0.140, y, 1.420, 0.580, a, 11, K["dk"], True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 1.850, y, COL_W - 1.850, 0.580, b, 10.6, INK, spacing=1.25, anchor=MSO_ANCHOR.MIDDLE)

cap(s, COL_R, 3.320, COL_W, "이 문서들이 실제로 한 일", K["pri"])
for i, (a, b) in enumerate([("의사결정을 남긴다", "결정이 번복되면 이유와 날짜를 함께 기록해 같은 논쟁을 반복하지 않도록"),
                            ("등급으로 관리한다", "오픈을 막는 P0와 오픈 전 정리 P1을 분리해 일정 판단의 근거로"),
                            ("회사 간 계약이 된다", "필드 매핑표가 세 개발사의 개발 범위와 책임 경계를 정하는 기준 문서로")]):
    y = 3.680 + i * 0.880
    rr(s, COL_R, y, COL_W, 0.780, WHITE, LINE)
    text(s, COL_R + 0.240, y + 0.090, COL_W - 0.480, 0.280, a, 11.5, INK, True)
    text(s, COL_R + 0.240, y + 0.380, COL_W - 0.480, 0.340, b, 10.2, SUB, spacing=1.25)
text(s, ML, 6.720, CW, 0.260,
     "협력사명은 자사몰 개발사 · ERP 개발사 · 물류 개발사 · 통관대행사 · 본인확인기관으로 치환한 사본을 사용했고 수치·계정 정보는 마스킹했습니다.",
     9.6, MUTE)

# ═════════════════════════════════════════ 08 검증과 확장
s = blank(p); head(s, "05", "검증과 확장", "06 / 06", K)
cap(s, COL_L, BODY, COL_W, "운영 환경 30건 · 테스트 목적 30/30 통과", K["pri"])
for i, (a, b) in enumerate([("정상 건", "정상 접수 성공"),
                            ("의도 오류 건", "예상한 오류를 정확히 반환"),
                            ("오류 수정 건", "원천 수정 후 새 제출번호로 재제출 성공")]):
    y = BODY + 0.330 + i * 0.520
    rr(s, COL_L, y, COL_W, 0.450, K["bg"] if i % 2 == 0 else WHITE, LINE)
    text(s, COL_L + 0.220, y, 1.600, 0.450, a, 11.3, K["dk"], True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 1.900, y, COL_W - 2.100, 0.450, b, 10.8, INK, anchor=MSO_ANCHOR.MIDDLE)
cap(s, COL_L, 3.500, COL_W, "핵심 E2E 시나리오", K["pri"])
for i, t in enumerate(["간소화 · 일반 본인확인 분기", "한 주문의 다중 운송장 분할 제출",
                       "운송장 미배정 · 부분배정 상태의 제출 차단",
                       "국내 통관 데이터와 해외 물류 데이터 병합 · 중복방지",
                       "자사몰 · 외부 채널별 데이터 프로필"]):
    y = 3.840 + i * 0.450
    rr(s, COL_L, y, COL_W, 0.380, WHITE, LINE)
    text(s, COL_L + 0.220, y, COL_W - 1.300, 0.380, t, 10.6, INK, anchor=MSO_ANCHOR.MIDDLE)
    label(s, COL_L + COL_W - 1.020, y + 0.045, 0.800, 0.290, "PASS", 9.5, GRN, GRN_BG, radius=0.14)

cap(s, COL_R, BODY, COL_W, "오류 복구 루프", K["pri"])
for i, (t, n) in enumerate([("제출", ""), ("오류 코드 회신", "약 1분"),
                            ("관리자 화면에서 원인 확인", "어떤 값이 왜 틀렸는지"),
                            ("원천 구분 후 수정", "자사 / 해외"), ("새 제출번호로 재제출", "처리창 내")]):
    y = BODY + 0.330 + i * 0.490
    rr(s, COL_R, y, 3.500, 0.390, K["pri"] if i == 0 else K["bg"])
    text(s, COL_R + 0.200, y, 3.100, 0.390, t, 11, WHITE if i == 0 else K["dk"], True, anchor=MSO_ANCHOR.MIDDLE)
    if n: text(s, COL_R + 3.620, y, 2.210, 0.390, n, 10, SUB, anchor=MSO_ANCHOR.MIDDLE)
    if i < 4:
        text(s, COL_R + 0.100, y + 0.380, 0.300, 0.110, "▼", 7.5, K["mid"], False, PP_ALIGN.CENTER)
rr(s, COL_R, 4.290, COL_W, 0.660, RED)
text(s, COL_R + 0.240, 4.290, COL_W - 0.480, 0.660,
     "11:50까지 해결되지 않으면 수동 통관 목록으로 격리\n자동 재제출과 자동 출고를 차단",
     11.2, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.32)

cap(s, COL_R, 5.150, COL_W, "조율에서 실제로 뚫은 것", K["pri"])
for i, (who, fix) in enumerate([
        ("ERP 개발사", "세관 제출용과 정보 조회용 로직이 따로 돌아 조회가 느렸습니다.\n두 로직을 잇는 중간 API를 새로 정의해 경로를 하나로 만들었습니다."),
        ("통관대행사", "API 필수값이 제출 기준에만 맞춰져 취합 로직과 맞지 않았습니다.\nERP용과 자사몰용 API를 각각 만들도록 요청해 분리했습니다.")]):
    y = 5.490 + i * 0.720
    label(s, COL_R, y + 0.060, 1.250, 0.300, who, 10, WHITE, K["pri"])
    text(s, COL_R + 1.400, y, COL_W - 1.400, 0.640, fix, 10.2, SUB, spacing=1.30)

# ═════════════════════════════════════════ 09 B2B 표지
s = blank(p)
cover(s, B, "01 / 06", "Project 02.   운영 디지털화 — 수기 운영을 설정형 시스템으로",
      "B2B", "회원제 폐쇄몰",
      "산후조리원 전용 B2B 폐쇄몰 구축",
      ["계약마다 달라지는 상품·단가·결제·포인트 규칙을",
       "운영자가 직접 설정하는 정책 구조로 바꿨습니다."],
      [(2.990, "Concept", 3.290, ["계약 예외를 코드가 아니라 관리자 설정값으로 분해"]),
       (3.650, "Main Target", 3.900, ["계약 조리원 담당자 · 내부 운영팀 · 공급사 5곳"]),
       (4.310, "Project Goal", 4.560, ["오출고 제거 · 입금 대조 자동화 · 신규 계약을 개발 없이 처리"])],
      "30개소에서 되던 것이 140개소에서 깨졌다",
      [("매일 입금 대조", "15~30분"), ("월말 정산", "2명 × 약 7시간"),
       ("오출고", "월 2~3건"), ("발주량 취합", "수기 시트")],
      [("기간", "2026.06 기획 → 07.02 개발 착수 → 07.16 핵심 서비스 전체 가동 → 08.06~08.13 운영 고도화·QA → 08.31 최종 검수"),
       ("Role", "문제 정의 / 플랫폼 대안 검토 / 서비스 정책·요구사항 정의 / 외주 개발 조율 / QA·UAT / 운영 정책 설계"),
       ("협업", "커머스 솔루션 구축사 · 내부 운영팀 · 영업 · 공급사 5곳"),
       ("산출물", "요구사항 대장 · 과업지시서 · 4계정 정책 매트릭스 · 권한표 · UAT 시나리오 · 운영 매뉴얼"),
       ("규모", "계약 조리원 약 140개소 · 물품 지원 계약 약 77개소 · 공급사 5곳")])

# ═════════════════════════════════════════ 10 배경과 AS-IS
s = blank(p); head(s, "01", "배경과 AS-IS 업무 흐름", "02 / 06", B)
cap(s, ML, BODY, 7.000, "AS-IS — 여섯 단계가 모두 사람 손을 거쳤다", B["pri"])
SW = (CW - 0.300) / 6
for i, (a, b) in enumerate([("구글 폼 주문", "조리원이 직접 작성"), ("시트 취합", "운영팀 수기"),
                            ("입금 대조", "계좌 내역 대조"), ("거래명세서 메일", "건별 발송"),
                            ("ERP 등록", "재입력"), ("월말 정산", "수기 대조")]):
    x = ML + i * (SW + 0.060)
    rr(s, x, BODY + 0.330, SW, 0.680, WHITE, LINE)
    text(s, x + 0.080, BODY + 0.400, SW - 0.160, 0.280, a, 11, INK, True, PP_ALIGN.CENTER)
    text(s, x + 0.080, BODY + 0.690, SW - 0.160, 0.260, b, 9.2, SUB, False, PP_ALIGN.CENTER)
    if i < 5:
        text(s, x + SW, BODY + 0.330, 0.060, 0.680, "›", 12, B["mid"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
text(s, ML, BODY + 1.090, CW, 0.260,
     "단계마다 기록되는 파일이 달라 같은 주문을 여러 번 옮겨 적었고, 계약 조건은 담당자 기억에 의존했습니다.", 10.8, SUB)

cap(s, COL_L, 3.320, COL_W, "수기 운영의 실제 부담", B["pri"])
for i, (a, v, d) in enumerate([("매일 입금 대조", "15 ~ 30분", "계좌 내역과 주문을 눈으로 맞춤"),
                               ("월말 정산", "2명 × 약 7시간", "주문·입금·정산 자료를 수기로 대조"),
                               ("오출고", "월 2 ~ 3건", "담당자 기억으로 계약 조건 판단")]):
    y = 3.680 + i * 0.640
    rect(s, COL_L, y, COL_W, 0.550, SOFT)
    text(s, COL_L + 0.220, y, 1.700, 0.550, a, 11.3, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 1.980, y, 1.400, 0.550, v, 12.5, B["pri"], True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 3.450, y, COL_W - 3.600, 0.550, d, 10, SUB, anchor=MSO_ANCHOR.MIDDLE)
rr(s, COL_L, 5.640, COL_W, 0.500, B["bg"])
text(s, COL_L + 0.220, 5.640, COL_W - 0.440, 0.500,
     "가장 먼저 막아야 할 것을 오출고와 입금 대조로 정했습니다.", 11.5, B["dk"], True, anchor=MSO_ANCHOR.MIDDLE)
text(s, COL_L, 6.280, COL_W, 0.520,
     "빈도가 높고 팀원의 시간을 매일 가져가고 있었기 때문입니다.\n실수 한 건이 곧 재출고 비용과 신뢰 손실로 이어졌습니다.", 10.2, SUB, spacing=1.32)

cap(s, COL_R, 3.320, COL_W, "문제를 어떻게 정의했나", B["pri"])
rect(s, COL_R, 3.680, COL_W, 1.020, B["bg"])
text(s, COL_R + 0.240, 3.680, COL_W - 0.480, 1.020,
     "주문 접수 화면이 불편한 것이 아니라,\n계약·결제·주문 상태를 사람이 매번 다시 판단하는\n운영 구조가 문제였습니다.",
     12.3, B["dk"], True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.38)
for i, t in enumerate(["계약 조건이 계약서와 담당자 기억에만 있어 시스템이 검증할 수 없었습니다.",
                       "결제 방식이 다른 상품이 한 주문에 섞여도 막을 방법이 없었습니다.",
                       "주문 데이터가 폼·시트·메일·ERP에 흩어져 어디가 원본인지 불명확했습니다."]):
    y = 4.840 + i * 0.450
    rect(s, COL_R, y + 0.135, 0.100, 0.100, B["mid"], shape=MSO_SHAPE.OVAL)
    text(s, COL_R + 0.240, y, COL_W - 0.240, 0.390, t, 11, INK, spacing=1.25)
text(s, COL_R, 6.280, COL_W, 0.520,
     "사람의 판단에 의존하는 구조를\n설정으로 검증되는 구조로 바꾸는 것이 목표였습니다.", 10.8, B["dk"], True, spacing=1.32)

# ═════════════════════════════════════════ 11 대안 비교와 오픈 범위
s = blank(p); head(s, "02", "대안 비교와 오픈 범위", "03 / 06", B)
cap(s, ML, BODY, 5.000, "검토한 세 가지 방식", B["pri"])
ry = table(s, ML, BODY + 0.310,
           [("대안", 3.000), ("장점", 2.300), ("한계", 2.800), ("판단 근거", 2.500), ("결정", 1.395)],
           [["맞춤형 웹서비스 직접 개발", "요구사항을 그대로 구현 가능",
             "서버 운영과 유지보수까지 직접 책임 · 일정과 비용 부담", "유지보수 책임", "제외"],
            ["타 쇼핑몰 솔루션 도입", "구축 속도가 빠름",
             "운영팀이 새 환경을 다시 학습해야 함 · 자사몰과 이원화", "운영팀 러닝커브", "제외"],
            ["기존 커머스 솔루션 + 정책 커스텀", "기존 자사몰과 동일 기반 · 러닝커브 최소",
             "기본 기능 밖 정책은 별도 개발이 필요",
             "러닝커브·인프라 충족\n유지보수 분담 · 예산 내", "채택"]],
           rowh=0.660, hi=2, pal=B, sz=10.8)
rect(s, ML, ry + 0.020, CW, 0.400, SOFT)
text(s, ML + 0.240, ry + 0.020, CW - 0.480, 0.400,
     "결정 기준 순서   ①  운영팀 러닝커브      ②  기존 인프라 일치      ③  유지보수 책임 범위      ④  구축 비용",
     11, INK, anchor=MSO_ANCHOR.MIDDLE)

cap(s, COL_L, 4.900, COL_W, "오픈에 반드시 필요했던 다섯 가지", B["pri"])
MW = (COL_W - 0.240) / 5
for i, t in enumerate(["계정별 상품 분리", "결제 방식 제어", "입금 자동 매칭", "공급사 권한", "주문 DB화"]):
    x = COL_L + i * (MW + 0.060)
    rr(s, x, 5.240, MW, 0.760, WHITE, B["mid"])
    text(s, x + 0.050, 5.300, MW - 0.100, 0.280, "0%d" % (i + 1), 10, B["pri"], True, PP_ALIGN.CENTER)
    text(s, x, 5.580, MW, 0.300, t, 9.2, INK, True, PP_ALIGN.CENTER, spacing=1.15)
text(s, COL_L, 6.100, COL_W, 0.260, "이 다섯 가지가 안 되면 오픈하지 않는다는 기준으로 범위를 관리했습니다.", 10.5, SUB)

cap(s, COL_R, 4.900, COL_W, "범위에서 뺀 것과 그 이유", RED)
for i, (a, b) in enumerate([("PG 결제", "마진이 낮아 수수료 부담이 커서 주 결제수단에서 제외"),
                            ("견적서 신규 개발", "예산을 초과해 솔루션 기본 거래명세서로 대체"),
                            ("포인트 완전 자동화", "계약 상태 변경 시 지급 번복 위험이 있어 관리자 확인 후 실행")]):
    y = 5.240 + i * 0.560
    label(s, COL_R, y + 0.060, 1.700, 0.360, a, 10, GOLD, GOLD_BG, radius=0.06)
    text(s, COL_R + 1.840, y, COL_W - 1.840, 0.480, b, 10.3, SUB, spacing=1.25, anchor=MSO_ANCHOR.MIDDLE)

# ═════════════════════════════════════════ 12 정책 매트릭스와 권한
s = blank(p); head(s, "03", "정책 매트릭스와 권한 설계", "04 / 06", B)
cap(s, ML, BODY, 5.000, "네 가지 계정 유형과 설정값", B["pri"])
ry = table(s, ML, BODY + 0.310,
           [("계정 유형", 2.000), ("노출 상품", 1.500), ("결제 방식", 1.700), ("포인트", 1.500), ("공급사 노출", 1.400)],
           [["분유 (선불)", "분유 SKU", "선불 · 무통장", "지원 포인트", "분유 공급사"],
            ["분유 (후불)", "분유 SKU", "후불 · 월 정산", "지원 포인트", "분유 공급사"],
            ["물품 (선불)", "물품 SKU", "선불 · 무통장", "미적용", "물품 공급사"],
            ["물품 (후불)", "물품 SKU", "후불 · 월 정산", "미적용", "물품 공급사"]],
           rowh=0.470, zebra=B["bg2"], pal=B, sz=10.8)
rect(s, ML, ry + 0.030, 8.100, 0.700, B["bg"])
text(s, ML + 0.240, ry + 0.030, 7.620, 0.700,
     "유형을 코드에 고정한 것이 아닙니다. 영업매니저가 계약 조건을 전달하면 운영팀이 관리자 화면에서\n상품·단가·결제·포인트·공급사 범위를 직접 설정합니다. 새 계약이 생겨도 개발이 필요 없습니다.",
     11, INK, anchor=MSO_ANCHOR.MIDDLE, spacing=1.35)

cap(s, 8.900, BODY, 3.762, "공급사에 연 것과 닫은 것", B["pri"])
rr(s, 8.900, BODY + 0.310, 3.762, 1.040, B["bg"])
text(s, 9.140, BODY + 0.400, 3.300, 0.250, "허용", 11, B["dk"], True)
text(s, 9.140, BODY + 0.670, 3.300, 0.600,
     "배정된 주문 조회 · 목록 다운로드\n주문 상태 변경 · 송장번호 입력", 10.3, INK, spacing=1.32)
rr(s, 8.900, BODY + 1.440, 3.762, 1.340, "FCEAEA")
text(s, 9.140, BODY + 1.530, 3.300, 0.250, "차단", 11, RED, True)
text(s, 9.140, BODY + 1.800, 3.300, 0.900,
     "계약 단가 · 정산 정보\n타 공급사 주문\n고객 개인정보", 10.3, INK, spacing=1.32)
rect(s, 8.900, BODY + 2.870, 3.762, 0.620, B["dk"])
text(s, 9.100, BODY + 2.870, 3.362, 0.620,
     "결제 방식이 다른 상품의\n혼합 주문은 양방향 차단", 11.2, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.30)

cap(s, ML, 5.560, CW, "하나의 주문 DB를 세 주체가 다른 범위로 본다", B["pri"])
for i, (a, b) in enumerate([("조리원 담당자", "내 주문 · 내 단가 · 거래명세서"),
                            ("내부 운영팀", "전체 주문 · 입금 매칭 · 정산 · 포인트"),
                            ("공급사", "배정된 주문 · 송장 입력")]):
    x = ML + i * 4.070
    rr(s, x, 5.900, 3.860, 0.660, WHITE, LINE)
    text(s, x + 0.220, 5.960, 3.420, 0.280, a, 11.3, B["dk"], True)
    text(s, x + 0.220, 6.250, 3.420, 0.260, b, 9.8, SUB)
text(s, ML, 6.720, CW, 0.260,
     "이전에는 폼·시트·메일·ERP에 흩어져 있어 무엇이 원본인지 매번 확인해야 했습니다.", 10.2, MUTE)

# ═════════════════════════════════════════ 13 사용자 여정과 화면
s = blank(p); head(s, "04", "사용자 여정과 화면", "05 / 06", B)
JW, JX = 2.130, 1.860
for i, t in enumerate(["주문", "결제 · 입금", "확인", "출고", "정산"]):
    rect(s, JX + i * (JW + 0.060), BODY, JW, 0.380, B["dk"])
    text(s, JX + i * (JW + 0.060), BODY, JW, 0.380, t, 11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
for li, (who, cells) in enumerate([
    ("조리원 담당자", ["계정 로그인 후\n계약 상품만 노출", "선불 무통장 입금\n또는 후불 확정", "주문 상태 확인",
                  "배송 안내 문자 수신", "거래명세서\n목록에서 즉시 다운로드"]),
    ("내부 운영팀", ["계약 조건대로\n계정·단가 설정", "입금 자동 매칭\n미매칭만 수기 확인", "혼합주문 차단 확인",
                 "공급사 배정", "미발주 계약처 자동 제외\n포인트 확인 후 실행"]),
    ("공급사", ["", "", "배정 주문 조회·다운로드", "송장번호 직접 입력", ""])]):
    y = BODY + 0.470 + li * 0.940
    rect(s, ML, y, 1.130, 0.860, B["bg"])
    text(s, ML, y, 1.130, 0.860, who, 10.3, B["dk"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.2)
    for ci, cl in enumerate(cells):
        if not cl: continue
        x = JX + ci * (JW + 0.060)
        rr(s, x, y, JW, 0.860, WHITE, LINE)
        text(s, x + 0.100, y, JW - 0.200, 0.860, cl, 9.6, INK, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.25)
text(s, ML, 4.900, CW, 0.260,
     "여정을 세 레인으로 나눠보니 같은 주문이 확인·출고 구간에서 세 주체를 동시에 거친다는 것이 드러났습니다. 이 구간에 혼합주문 차단과 공급사 권한을 배치했습니다.",
     10.5, SUB)

cap(s, ML, 5.320, CW, "실제 운영 화면", B["pri"])
for i, (a, b) in enumerate([("관리자 설정 화면", "회원등급 · 상품 노출 · 계약 단가"),
                            ("주문 목록 · 거래명세서", "목록에서 즉시 다운로드"),
                            ("혼합주문 차단", "결제 방식이 다른 상품 혼합 시 안내"),
                            ("공급사 전용 화면", "배정 주문 조회 · 송장 입력")]):
    x = ML + i * 3.050
    rr(s, x, 5.660, 2.870, 0.880, SOFT, LINE)
    text(s, x + 0.200, 5.760, 2.470, 0.280, a, 10.8, INK, True)
    text(s, x + 0.200, 6.050, 2.470, 0.400, b, 9.6, SUB, spacing=1.22)
text(s, ML, 6.660, CW, 0.240,
     "네 화면 모두 실제 운영 화면이며 고객·단가·주문정보는 마스킹했습니다.", 9.6, MUTE)

# ═════════════════════════════════════════ 14 검증과 결과
s = blank(p); head(s, "05", "검증과 결과", "06 / 06", B)
cap(s, COL_L, BODY, COL_W, "오픈 전 반드시 통과시킨 시나리오", B["pri"])
for i, t in enumerate(["계정 유형별 정상 주문 · 상품 노출 분기", "실제 입금 후 자동 매칭과 주문 상태 변경",
                       "결제 방식이 다른 상품의 혼합주문 차단", "공급사 송장 입력 후 상태 변경과 안내 문자 발송",
                       "후불 주문의 미입금 오인 방지", "미발주 계약처의 정산 대상 제외"]):
    y = BODY + 0.330 + i * 0.430
    rr(s, COL_L, y, COL_W, 0.370, WHITE, LINE)
    text(s, COL_L + 0.200, y, COL_W - 1.250, 0.370, t, 10.4, INK, anchor=MSO_ANCHOR.MIDDLE)
    label(s, COL_L + COL_W - 0.990, y + 0.045, 0.780, 0.280, "PASS", 9.3, GRN, GRN_BG, radius=0.14)
text(s, COL_L, 4.520, COL_W, 0.260, "무통장 자동 매칭과 안내 문자는 현재 실제 운영 중입니다.", 10.5, B["dk"], True)

cap(s, COL_R, BODY, COL_W, "도입 후 변화", B["pri"])
for i, (nm, bf, af, note) in enumerate([
    ("오출고", "월 2~3건", "0건", "직전 6개월 월평균 대비 · 07-16 전체 가동 후 약 1개월 · 원천 운영팀 출고·CS 이력"),
    ("월말 정산", "2명 × 약 7시간", "1명 × 30~60분", "경과 420분 · 총 14인시 · 원천 기존 수행자인 경영관리팀·운영팀 확인"),
    ("배송·주문·입금 문의", "2주 32건", "2주 19건", "2026.01~06 인입의 2주 환산 평균과 7~8월 배포 구간 2주 비교 · 초기 관찰값")]):
    y = BODY + 0.330 + i * 1.150
    rr(s, COL_R, y, COL_W, 1.040, B["bg"] if i % 2 == 0 else WHITE, LINE)
    text(s, COL_R + 0.220, y + 0.080, 2.100, 0.300, nm, 11.5, B["dk"], True)
    text(s, COL_R + 0.220, y + 0.390, 2.100, 0.540, bf, 10.2, SUB, spacing=1.25)
    text(s, COL_R + 2.420, y + 0.290, 0.320, 0.320, "→", 12, B["mid"], True, PP_ALIGN.CENTER)
    text(s, COL_R + 2.800, y + 0.210, 1.600, 0.420, af, 15, B["pri"], True)
    text(s, COL_R + 0.220, y + 0.730, COL_W - 0.440, 0.280, note, 8.6, MUTE, spacing=1.20)

cap(s, COL_L, 5.000, COL_W, "운영에서 실제로 없어진 일", B["pri"])
for i, (a, b) in enumerate([("영업매니저의 수기 발주량 시트", "DB에서 바로 조회"),
                            ("주문 없는 계약처 수기 제외", "정산에서 자동 제외"),
                            ("공급사 송장 중간 전달", "공급사가 직접 입력")]):
    y = 5.340 + i * 0.480
    text(s, COL_L, y, 3.100, 0.400, a, 10.2, SUB, anchor=MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 3.150, y, 0.300, 0.400, "→", 10, B["mid"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, COL_L + 3.520, y, 2.310, 0.400, b, 10.8, B["dk"], True, anchor=MSO_ANCHOR.MIDDLE)
text(s, ML, 6.900, CW, 0.280,
     "계약상 기본 개발요구 22건은 08-31 최종 검수를 기준으로 운영 QA를 이어갔고, 08-13 시점 17건이 운영 서버 확인 전 상태였습니다.",
     9.8, MUTE)

# ═════════════════════════════════════════ 15 BIGSEE 표지
s = blank(p)
cover(s, M, "01 / 07", "Project 03.   제품 기획 — 현장에서 본 병목을 서비스로",
      "BIG", "마케팅 리서치 SaaS",
      "채널 조사를 매일 도는 제품으로",
      ["마케팅은 결국 소스를 얼마나 빨리 보느냐의 문제였습니다.",
       "기존에 있던 크롤링 로직을 고도화해 매일 도는 형태로 바꿨습니다."],
      [(2.990, "Problem", 3.290, ["조사를 반복하는 사람은 채널마다 흩어진 숫자를 손으로 합치고",
                                  "그 근거는 다음 조사로 이어지지 않는다"]),
       (3.880, "Target 가설", 4.130, ["조사가 업무인 직무 — 광고대행사 실무자 (1순위 검증 가설, 확정 아님)"]),
       (4.540, "Goal", 4.790, ["판매 대상 정의 · 이탈 지점 특정 · 판매 범위(P0) 확정 · 검증 기준 정의"])],
      "어떻게 시작됐나",
      [("현장에서 병목 관찰", "SCM 실무 중"), ("크롤링 로직 고도화", "기존 로직 활용"),
       ("동료와 구체화", "브레인스토밍"), ("별도 서비스 개설", "현재 운영 중")],
      [("기간", "진행 중 · https://market-bigsee.com/"),
       ("Role", "아이디어 발의 · 제품·사업 기획 — 개발 구현은 동료 개발자가 전담"),
       ("담당", "문제 정의 / 기회 정의 / 타깃·페르소나 / 사용자 여정 / 제품 흐름 / AI 응답 정책 / 검증 기준")])

# ═════════════════════════════════════════ 16 문제 정의
s = blank(p); head(s, "01", "문제 정의", "02 / 07", M)
cap(s, ML, BODY, 7.000, "AS-IS — 마케터가 키워드 하나를 조사하는 방식", M["pri"])
AW = (CW - 0.240) / 5
for i, t in enumerate(["네이버 검색광고에서\n검색량 · 경쟁도 확인", "데이터랩에서\n관심도 추이 따로 확인",
                       "유튜브 · 인스타를\n눈으로 훑어 반응 가늠", "엑셀에 숫자를\n손으로 옮겨 붙임",
                       "캡처를 모아\n보고서 작성"]):
    x = ML + i * (AW + 0.060)
    rr(s, x, BODY + 0.330, AW, 0.880, WHITE, LINE)
    rect(s, x + 0.180, BODY + 0.450, 0.300, 0.300, SOFT, shape=MSO_SHAPE.OVAL)
    text(s, x + 0.180, BODY + 0.450, 0.300, 0.300, str(i + 1), 10, SUB, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x + 0.180, BODY + 0.790, AW - 0.360, 0.380, t, 9.8, INK, spacing=1.20)
text(s, ML, BODY + 1.290, CW, 0.260,
     "다섯 개의 창을 열고 기준이 서로 다른 숫자를 손으로 합칩니다. 같은 키워드를 다음 달에 다시 조사하면 이 과정을 처음부터 반복합니다.", 10.8, SUB)

cap(s, ML, 3.500, CW, "그래서 생기는 문제", M["pri"])
for i, (k, t, d) in enumerate([("시간", "조사 한 건에 반나절", "채널마다 화면과 지표 기준이 달라 매번 다시 익혀야 합니다."),
                               ("근거 소실", "출처와 수집일이 안 남음", "보고서에 숫자만 남고 어디서 언제 가져왔는지가 사라집니다."),
                               ("비교 불가", "다음 조사와 이어지지 않음", "지난달과 같은 기준으로 비교할 수 없어 추세를 못 읽습니다.")]):
    x = ML + i * 4.070
    rect(s, x, 3.840, 3.860, 1.180, M["bg"])
    text(s, x + 0.260, 3.940, 3.340, 0.240, k, 9.8, M["pri"], True)
    text(s, x + 0.260, 4.220, 3.340, 0.300, t, 13.5, INK, True)
    text(s, x + 0.260, 4.580, 3.340, 0.380, d, 9.8, SUB, spacing=1.25)

cap(s, ML, 5.360, CW, "Problem Statement", M["pri"])
rect(s, ML, 5.700, CW, 0.900, INK)
text(s, ML + 0.320, 5.700, CW - 0.640, 0.900,
     "조사를 반복하는 사람은 채널마다 흩어진 숫자를 손으로 합치느라 반나절을 쓰고,\n그렇게 만든 근거는 다음 조사에 이어지지 않는다.",
     13.5, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.40)

# ═════════════════════════════════════════ 17 데스크 리서치
s = blank(p); head(s, "02", "데스크 리서치", "03 / 07", M)
cap(s, COL_L, BODY, COL_W, "시장 — 누가 이 조사를 반복하는가", M["pri"])
for i, (a, v, d) in enumerate([("국내 디지털 광고 시장", "성장 지속", "광고비가 늘수록 사전 조사 수요도 함께 늘어납니다"),
                               ("국내 광고회사", "다수 · 소규모 중심", "조사 전담 인력이 없는 곳일수록 도구 의존도가 높습니다"),
                               ("마케팅 리서치 도구 시장", "글로벌 성장세", "국내는 채널별 단일 도구가 대부분입니다")]):
    y = BODY + 0.330 + i * 0.740
    rr(s, COL_L, y, COL_W, 0.640, WHITE, LINE)
    text(s, COL_L + 0.240, y + 0.080, 3.200, 0.260, a, 11.3, INK, True)
    text(s, COL_L + 3.500, y + 0.080, 2.100, 0.260, v, 11, M["pri"], True, PP_ALIGN.RIGHT)
    text(s, COL_L + 0.240, y + 0.360, COL_W - 0.480, 0.240, d, 9.6, SUB)
rect(s, COL_L, 4.140, COL_W, 0.800, SOFT)
text(s, COL_L + 0.240, 4.140, COL_W - 0.480, 0.800,
     "리서치에서 세운 첫 가설\n채널을 많이 조사하는 사람일수록 통합의 가치를 크게 느낀다. 따라서 조사가 잦은 직무에 먼저 팔린다.",
     10.8, INK, spacing=1.38, anchor=MSO_ANCHOR.MIDDLE)

cap(s, COL_R, BODY, COL_W, "도구 지형 — 국내는 두 갈래로 갈려 있다", M["pri"])
ry = table(s, COL_R, BODY + 0.330,
           [("유형", 1.600), ("대표 도구", 1.900), ("다루는 범위", 1.500), ("결론까지", 0.830)],
           [["검색량 도구", "블랙키위 · 키워드마스터", "네이버 검색 중심", "미제공"],
            ["소셜 반응 도구", "썸트렌드 계열", "SNS 언급 · 감성", "미제공"],
            ["글로벌 SEO 도구", "Ahrefs · SEMrush 계열", "구글 중심", "일부"],
            ["BIGSEE", "5개 채널 통합", "검색 + 반응 + 뉴스", "AI 해설"]],
           rowh=0.520, hi=3, pal=M, sz=10.3)
rect(s, COL_R, ry + 0.040, COL_W, 0.840, M["bg"])
text(s, COL_R + 0.240, ry + 0.040, COL_W - 0.480, 0.840,
     "국내 도구는 검색량 아니면 SNS 반응 한쪽만 봅니다.\n둘을 함께 보고 그래서 무엇을 할지까지 남기는 화면이 없었습니다.",
     10.8, M["dk"], spacing=1.38, anchor=MSO_ANCHOR.MIDDLE)

cap(s, ML, 5.320, CW, "이 리서치로 정해진 것", M["pri"])
for i, (n, t, d) in enumerate([
        ("01", "먼저 팔 대상을 조사 빈도로 좁힌다",
         "시장 크기가 아니라 조사를 반복하는 빈도가\n통합 도구의 가치를 결정한다고 봤습니다."),
        ("02", "채널 추가가 아니라 결론 제공으로 간다",
         "국내 도구가 이미 채널별로 존재하므로\n같은 축에서 경쟁하지 않기로 했습니다."),
        ("03", "경쟁 비교는 미검증으로 남긴다",
         "공개 기능 페이지만 봤고 유료 플랜을\n전수 비교하지 않았습니다.")]):
    x = ML + i * 4.070
    rr(s, x, 5.660, 3.860, 1.060, WHITE, LINE)
    text(s, x + 0.240, 5.750, 0.400, 0.240, n, 10, M["pri"], True)
    text(s, x + 0.700, 5.750, 2.900, 0.240, t, 11.3, INK, True)
    text(s, x + 0.240, 6.070, 3.380, 0.560, d, 9.8, SUB, spacing=1.30)

text(s, ML, 6.900, CW, 0.280,
     "출처 · DMC리포트 · KOBACO 2025 디지털 광고 시장 보고서 / 한국광고총연합회 2025 광고회사 현황조사 / IMARC 마케팅 리서치 시장 자료",
     8.8, MUTE)

# ═════════════════════════════════════════ 18 포지셔닝
s = blank(p); head(s, "03", "포지셔닝", "04 / 07", M)
MX, MY, MW2, MH = ML, BODY + 0.200, 8.100, 4.400
rect(s, MX, MY, MW2, MH, SOFT)
hline(s, MX, MY + MH / 2, MW2, "C8C8CE")
rect(s, MX + MW2 / 2, MY, 0.014, MH, "C8C8CE")
text(s, MX, MY - 0.300, MW2, 0.240, "실행까지 연결", 10, SUB, False, PP_ALIGN.CENTER)
text(s, MX, MY + MH + 0.070, MW2, 0.240, "데이터 제공에서 끝", 10, SUB, False, PP_ALIGN.CENTER)
text(s, MX + 0.060, MY + MH / 2 - 0.290, 1.400, 0.240, "단일 채널", 10, SUB)
text(s, MX + MW2 - 1.460, MY + MH / 2 - 0.290, 1.400, 0.240, "멀티 채널", 10, SUB, False, PP_ALIGN.RIGHT)
for x, y, w, h, t in [(MX + 0.400, MY + 2.720, 1.900, 0.520, "썸트렌드 계열"),
                      (MX + 0.400, MY + 3.400, 1.900, 0.520, "블랙키위 · 키워드마스터"),
                      (MX + 2.400, MY + 2.980, 1.650, 0.620, "Ahrefs · SEMrush\n(구글 중심)"),
                      (MX + 4.900, MY + 3.500, 1.700, 0.520, "범용 생성 AI")]:
    rr(s, x, y, w, h, WHITE, LINE)
    text(s, x + 0.100, y, w - 0.200, h, t, 10, SUB, False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.18)
rr(s, MX + 5.400, MY + 2.620, 2.200, 0.720, M["br"])
text(s, MX + 5.400, MY + 2.680, 2.200, 0.300, "BIGSEE 현재", 11.5, WHITE, True, PP_ALIGN.CENTER)
text(s, MX + 5.400, MY + 2.970, 2.200, 0.260, "5채널 통합 조회", 9.5, "FDEEE2", False, PP_ALIGN.CENTER)
rr(s, MX + 5.400, MY + 0.420, 2.200, 0.720, M["dk"])
text(s, MX + 5.400, MY + 0.480, 2.200, 0.300, "목표 위치 (P0)", 11.5, WHITE, True, PP_ALIGN.CENTER)
text(s, MX + 5.400, MY + 0.770, 2.200, 0.260, "조사 → 실행 → 성과", 9.5, "FDEEE2", False, PP_ALIGN.CENTER)
rect(s, MX + 6.470, MY + 1.180, 0.060, 1.400, M["pri"])
text(s, MX + 6.180, MY + 1.020, 0.640, 0.220, "▲", 10, M["pri"], True, PP_ALIGN.CENTER)

cap(s, 9.100, BODY + 0.200, 3.562, "읽는 법", M["pri"])
for i, (a, b) in enumerate([("가로축", "몇 개 채널을 한 번에 다루는가"),
                            ("세로축", "숫자를 주고 끝나는가, 다음 행동까지 남기는가"),
                            ("빈칸", "오른쪽 위 — 멀티채널이면서 실행까지 잇는 자리"),
                            ("현재", "BIGSEE는 오른쪽 아래. 통합은 됐지만 결론이 없다")]):
    y = BODY + 0.540 + i * 0.860
    rr(s, 9.100, y, 3.562, 0.760, WHITE, LINE)
    text(s, 9.320, y + 0.090, 3.120, 0.260, a, 10.5, M["pri"], True)
    text(s, 9.320, y + 0.370, 3.120, 0.320, b, 10, SUB, spacing=1.22)
rect(s, 9.100, BODY + 3.980, 3.562, 0.620, M["bg"])
text(s, 9.320, BODY + 3.980, 3.122, 0.620,
     "채널을 더 붙이는 대신\n세로축으로 올리기로 결정", 11, M["dk"], True, anchor=MSO_ANCHOR.MIDDLE, spacing=1.30)

# ═════════════════════════════════════════ 19 타깃과 페르소나
s = blank(p); head(s, "04", "타깃과 페르소나", "05 / 07", M)
cap(s, ML, BODY, CW, "타깃을 좁힌 기준", M["pri"])
for i, (a, b) in enumerate([("조사 빈도", "월 1회 이하는 제외. 주 단위로 조사하는 직무만."),
                            ("결정 권한", "조사 결과로 예산 · 일정을 실제로 움직이는가."),
                            ("대체 수단", "무료 도구로 버티는가, 이미 유료를 쓰는가.")]):
    x = ML + i * 4.070
    rect(s, x, BODY + 0.310, 3.860, 0.780, SOFT)
    text(s, x + 0.240, BODY + 0.400, 3.380, 0.260, a, 10.8, INK, True)
    text(s, x + 0.240, BODY + 0.680, 3.380, 0.340, b, 9.8, SUB, spacing=1.25)

cap(s, ML, 2.980, CW, "가설 페르소나 — 인터뷰 전 단계", M["pri"])
for i, (pid, rank, nm, freq, task, pain, want, hi_) in enumerate([
    ("P1", "1순위 검증 가설", "광고대행사 AE", "주 3~5건", "제안서에 넣을 근거를 매번 새로 만든다",
     "왜 이 키워드냐는 질문에 답할 출처가 없다", "출처 · 수집일이 찍힌 조사 근거와 제안서 초안", True),
    ("P2", "2순위 비교군", "인하우스 마케터", "주 1~2건", "캠페인 전 수요를 확인하고 광고비를 배분한다",
     "채널마다 기준이 달라 어디에 더 쓸지 근거가 약하다", "채널 간 비교가 같은 기준으로 되는 화면", False),
    ("P3", "3순위 비교군", "다채널 활성 셀러", "주 2~3건", "상품을 올리기 전 수요와 경쟁을 본다",
     "쇼핑 클릭과 검색량을 따로 보느라 판단이 늦다", "SKU 단위로 묶여 저장되는 조사 결과", False)]):
    x = ML + i * 4.070
    rr(s, x, 3.320, 3.860, 3.560, WHITE, LINE)
    rect(s, x, 3.320, 3.860, 0.060, M["pri"] if hi_ else "C8C8CE")
    label(s, x + 0.240, 3.500, 0.520, 0.300, pid, 9.5, WHITE if hi_ else SUB, M["pri"] if hi_ else SOFT, radius=0.06)
    text(s, x + 0.880, 3.500, 2.700, 0.300, rank, 9.5, SUB, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.240, 3.900, 3.380, 0.320, nm, 15, INK, True)
    hline(s, x + 0.240, 4.310, 3.380)
    for j, (k, v) in enumerate([("조사 빈도", freq), ("과업", task), ("페인 포인트", pain)]):
        yy = 4.430 + j * 0.660
        text(s, x + 0.240, yy, 3.380, 0.240, k, 9, MUTE, True)
        text(s, x + 0.240, yy + 0.240, 3.380, 0.400, v, 10, SUB, spacing=1.25)
    text(s, x + 0.240, 6.400, 3.380, 0.240, "기대하는 것", 9, MUTE, True)
    text(s, x + 0.240, 6.640, 3.380, 0.400, want, 10, M["pri"], True, spacing=1.25)
text(s, ML, 7.020, CW, 0.260,
     "누구나 쓰는 도구라고 소개하면 광고 문장을 쓸 대상이 없습니다. 한 직무를 정해야 한 문장이 나옵니다.", 10.2, SUB)

# ═════════════════════════════════════════ 20 사용자 여정 지도
s = blank(p); head(s, "05", "사용자 여정 지도", "06 / 07", M)
EW = (CW - 0.300) / 6
JSTEP = [("키워드 정하기", "무엇을 조사할지 결정", "사내 요청 · 기획서", "보통"),
         ("채널별 조회", "다섯 채널을 각각 확인", "각 플랫폼 · BIGSEE", "낮음"),
         ("결과 읽기", "검색량 · 반응 · 뉴스 비교", "BIGSEE 결과 화면", "높음"),
         ("내보내기", "CSV · HTML로 저장", "내보내기 버튼", "보통"),
         ("보고서 만들기", "엑셀 · PPT로 옮겨 붙임", "사내 문서 · 메일", "낮음"),
         ("실행 · 성과 확인", "광고 · 콘텐츠 집행 후 확인", "광고 관리자 · GA", "낮음")]
EMO = {"높음": 0.0, "보통": 0.230, "낮음": 0.460}
for i, (a, b, c, e) in enumerate(JSTEP):
    x = ML + i * (EW + 0.060)
    rect(s, x, BODY, EW, 0.400, M["bg"] if i % 2 == 0 else M["bg2"])
    text(s, x + 0.060, BODY, EW - 0.120, 0.400, a, 10.5, M["dk"], True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    text(s, x + 0.060, BODY + 0.440, EW - 0.120, 0.300, b, 9.4, INK, False, PP_ALIGN.CENTER, spacing=1.18)
    text(s, x + 0.060, BODY + 0.760, EW - 0.120, 0.240, c, 8.8, MUTE, False, PP_ALIGN.CENTER)
rect(s, ML, 2.700, CW, 1.180, SOFT)
text(s, ML + 0.180, 2.760, 2.000, 0.220, "감정선", 9, MUTE, True)
for i in range(len(JSTEP) - 1):                # 점을 잇는 감정선
    x0 = ML + i * (EW + 0.060) + EW / 2
    x1 = ML + (i + 1) * (EW + 0.060) + EW / 2
    y0 = 2.975 + EMO[JSTEP[i][3]]
    y1 = 2.975 + EMO[JSTEP[i + 1][3]]
    cn = s.shapes.add_connector(1, I(x0), I(y0), I(x1), I(y1))
    cn.line.color.rgb = C(M["mid"]); cn.line.width = Pt(1.5)
for i, (a, b, c, e) in enumerate(JSTEP):
    x = ML + i * (EW + 0.060)
    cy = 2.900 + EMO[e]
    rect(s, x + EW / 2 - 0.075, cy, 0.150, 0.150, M["br"], shape=MSO_SHAPE.OVAL)
    text(s, x, cy + 0.175, EW, 0.220, e, 9, SUB, False, PP_ALIGN.CENTER)

cap(s, ML, 4.020, CW, "각 단계의 문제", M["pri"])
for i, t in enumerate(["조사 목적이 문서로 안 남는다", "탭 다섯 개, 기준이 제각각", "무엇이 좋은 수치인지 모른다",
                       "등급에 따라 잠겨 있다", "여기서 서비스 밖으로 나간다", "조사와 성과가 안 이어진다"]):
    x = ML + i * (EW + 0.060)
    hot = i in (1, 4, 5)
    rr(s, x, 4.360, EW, 0.640, "FCEAEA" if hot else WHITE, LINE)
    text(s, x + 0.080, 4.360, EW - 0.160, 0.640, t, 9.4, RED if hot else INK,
         False, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.20)
cap(s, ML, 5.140, CW, "대응", M["pri"])
for i, (t, p0) in enumerate([("조사 브리프로 목적 고정", False), ("한 화면 통합 (현재 제공)", False),
                             ("AI 해설 · 벤치마크 (제공)", False), ("잠금 사유와 대안 안내", False),
                             ("근거 기반 초안 자동 생성 (P0)", True), ("캠페인 · UTM · CRM (P0)", True)]):
    x = ML + i * (EW + 0.060)
    rr(s, x, 5.480, EW, 0.640, M["br"] if p0 else M["bg"])
    text(s, x + 0.080, 5.480, EW - 0.160, 0.640, t, 9.4, WHITE if p0 else M["dk"],
         p0, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE, 1.20)
text(s, ML, 6.320, CW, 0.560,
     "감정이 두 번 꺾입니다. 두 번째 꺾임(보고서 · 실행)에서 사용자는 서비스를 떠나고, 그 뒤 무슨 일이 있었는지 제품은 알지 못합니다.\nP0 범위를 5 · 6단계에 몰아넣은 이유입니다.",
     10.5, INK, spacing=1.35)

# ═════════════════════════════════════════ 21 가설·검증 보드
s = blank(p); head(s, "06", "가설과 검증 보드", "07 / 07", M)
hcols = [("ID", 0.800), ("가설", 4.100), ("통과 기준", 3.100), ("현재 증거", 2.700), ("상태", 1.295)]
STC = {"미시작": (SUB, SOFT), "진행": (M["dk"], M["bg"]), "준비": (GOLD, GOLD_BG)}
cx = ML
for lb, w in hcols:
    text(s, cx, BODY, w, 0.240, lb, 9.8, SUB, True); cx += w
hline(s, ML, BODY + 0.300, CW, M["mid"])
ry = BODY + 0.400
for i, r in enumerate([
    ["H-001", "대행사는 멀티채널 제안서 조사 문제에 돈을 낸다", "파일럿 제안 5건 중 3건 이상 결제", "공개 사용사례만 · 직접 증거 없음", "미시작"],
    ["H-002", "조사 시간을 의미 있게 줄인다", "프로젝트당 중앙값 60분 이상 절감", "미측정", "미시작"],
    ["H-003", "유료 고객이 4주 동안 반복 사용한다", "유료 고객 60% 이상 4주차 활성", "미측정", "미시작"],
    ["H-004", "사용자별 직접원가가 매출의 25% 이하다", "2주 연속 직접원가 비율 25% 이하", "비용 항목만 파악", "미시작"],
    ["H-010", "기존 약 20명 안에서 첫 고객군을 찾을 수 있다", "동일 문제 6명 이상 · 데모 5건 중 유료 2건", "등록 약 20명이라는 전달만 있음", "진행"],
    ["H-012", "커뮤니티와 Meta 동시 시험으로 고객군을 비교한다", "한 페르소나에서 유료 2건", "채널 공개 근거만 · 실적 없음", "준비"],
    ["H-013", "대행사가 인하우스·셀러보다 확장 가치가 높다", "결제·반복·팀 신호의 종합 우위", "업무 빈도 가설만", "진행"]]):
    if i % 2 == 0: rect(s, ML, ry - 0.045, CW, 0.480, M["bg2"])
    cx = ML
    for j, (lb, w) in enumerate(hcols):
        if j == 4:
            fg, bgc = STC[r[4]]
            label(s, cx, ry + 0.060, 0.920, 0.290, r[4], 9.3, fg, bgc, radius=0.14)
        else:
            text(s, cx, ry, w - 0.150, 0.400, r[j], 10.3,
                 M["pri"] if j == 0 else (M["dk"] if j == 2 else (INK if j == 1 else SUB)),
                 j in (0, 2), anchor=MSO_ANCHOR.MIDDLE, spacing=1.20)
        cx += w
    ry += 0.480
hline(s, ML, ry - 0.045, CW, M["mid"])
text(s, ML, ry + 0.070, CW, 0.240, "전체 14건 중 7건 발췌 · 원본은 가설·검증 대장에 유지", 9.6, MUTE)

cap(s, COL_L, 5.700, 7.400, "검증 원칙", M["pri"])
for i, t in enumerate(["인터뷰의 호의적 반응보다 실제 결제·반복·게시를 강한 증거로 본다",
                       "최소 표본과 기간을 사전에 정하고, 한 번의 성공으로 통과시키지 않는다",
                       "목표 미달이어도 숨기지 않고 기각 · 수정 · 추가시험 중 하나로 판정한다"]):
    y = 6.040 + i * 0.400
    rect(s, COL_L, y + 0.130, 0.090, 0.090, M["mid"], shape=MSO_SHAPE.OVAL)
    text(s, COL_L + 0.220, y, 7.100, 0.340, t, 10.3, INK, anchor=MSO_ANCHOR.MIDDLE)
rect(s, 8.300, 5.700, 4.362, 1.540, M["dk"])
text(s, 8.540, 5.840, 3.900, 0.280, "실제로 수정한 가설", 11.5, WHITE, True)
text(s, 8.540, 6.180, 3.900, 0.900,
     "H-011 순차 집행을 폐기하고\nH-012 동시 비교로 대체했습니다.\n상태를 '수정'으로 남기고 지우지 않았습니다.",
     10.2, M["lt"], spacing=1.35)

# ═════════════════════════════════════════ 22 마무리
s = blank(p)
rect(s, 0, 0, 13.3333, 7.5, WHITE)
rect(s, 0, 0, 0.24, 7.5, N["pri"])
text(s, 10.462, 0.345, 2.200, 0.260, "22 / 22", 10.5, MUTE, False, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
text(s, 1.200, 1.050, 8.600, 0.900,
     "현장을 구조로 바꾸고, 이해관계를 조율하고,\n운영 안착까지 보는 PM입니다.", 26, INK, True, spacing=1.32)
cap(s, 1.200, 2.420, 6.000, "세 프로젝트를 지나며 실제로 바뀐 것", N["pri"])
for i, (t, d) in enumerate([
        ("계획을 세우고 들어가는 방식에서, 만들어 보고 고치는 방식으로",
         "예전에는 계획을 충분히 세운 뒤 착수했습니다. 지금은 프로토타입을 여러 번 돌리고 버전을 올립니다."),
        ("먼저 그리지 않고, 먼저 듣는다",
         "관세청 프로젝트를 다시 한다면 API 가이드를 놓고 전체 회의부터 열어 각 개발사의 고충을 먼저 파악하겠습니다.\n부분 정보만 오간 것이 일정 지연의 원인이었습니다."),
        ("진행 상황을 글이 아니라 그림으로 공유한다",
         "중간 공유가 부족하다는 지적을 받고, 요약과 시각화 중심으로 공유 형식을 바꿨습니다.")]):
    y = 2.780 + i * 1.180
    text(s, 1.200, y, 0.520, 0.300, "0%d" % (i + 1), 13, N["mid"], True)
    text(s, 1.820, y, 7.600, 0.300, t, 14, INK, True)
    text(s, 1.820, y + 0.360, 7.600, 0.660, d, 10.5, SUB, spacing=1.35)
hline(s, 1.200, 6.420, 8.600)
text(s, 1.200, 6.580, 1.100, 0.260, "함께 제출", 10.5, N["mid"], True)
text(s, 2.400, 6.580, 7.400, 0.260, "이력서 · 마켓빅시 AI 응답 정책 발췌", 11, INK)
text(s, 1.200, 6.880, 1.100, 0.260, "서비스", 10.5, N["mid"], True)
text(s, 2.400, 6.880, 7.400, 0.260, "https://market-bigsee.com/", 11, N["pri"])
text(s, 10.400, 6.580, 2.262, 0.560,
     "박종혁\nparkjonghyeok2000@gmail.com", 10, SUB, False, PP_ALIGN.RIGHT, spacing=1.35)

p.save(OUT)
print("saved:", OUT)
print("slides:", len(p.slides._sldIdLst))
